"""Render the current practice PDF and a separate teacher-template report."""
from pathlib import Path
import json
import re
import shutil
import subprocess
import sys
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw
import pymupdf
import win32com.client

ROOT = Path(__file__).resolve().parents[1]
PRACTICE = '20241328_蔡贸俊_实验二_4-6学时_实践记录'
REPORT = '20241328_蔡贸俊_实验二_密码算法实现'


def font(run, name='宋体', size=10.5):
    run.font.name, run.font.size = name, Pt(size)
    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), name)


def main():
    summary = json.loads((ROOT/'logs/crosscheck_results.json').read_text(encoding='utf-8'))
    assert summary['passed'] == 59
    first = Path(r'D:\HuaweiMoveData\Users\HUAWEI\Documents\Default Project\20241328-蔡贸俊-实验二-1-3学时')
    appendix = ROOT/'materials/第一阶段记录'
    appendix.mkdir(exist_ok=True)
    for name in ('20241328_蔡贸俊_实验二_1-3学时_实践记录.md', '20241328_蔡贸俊_实验二_1-3学时_实践记录.pdf', 'images/实际终端验证截图.png', 'logs/results.json', 'logs/verification.log', 'logs/git_log.txt'):
        dest = appendix/name
        dest.parent.mkdir(parents=True, exist_ok=True)
        if (first/name).is_file():
            shutil.copy2(first/name, dest)
    practice = (ROOT/(PRACTICE+'.md')).read_text(encoding='utf-8')
    combined = practice.replace('# 实验二：密码算法实现——4—6学时实践记录', '# 实验二：密码算法实现——综合实验报告', 1)
    note = '''### 1.3 第一阶段（1—3学时）已完成内容与两阶段关系

第一阶段于2026年10月22日在openEuler运行商用密码检测中心参考SM2/SM3/SM4代码，并用GmSSL交叉验证，36项断言通过。具体包括官方三算法自检、11组SM3输入、8组SM4双端加密与反向解密、SM2双向互操作和C3篡改拒绝；当时取得的标准及函数映射分析已完成。前半段实践记录、原始截图和日志副本保存在materials/第一阶段记录，仍对应其原始日期，没有把本次硬件结果写进旧日志。

第一阶段报告中记录了代码PDF恢复、MIRACL LP64配置、标准查询、编码互操作和效率测量；效率结果仅限其原固定参数与单次测量条件，不作为本次硬件性能对照。该阶段已发布课程仓库，最终文档提交08dec4448c78f7890e59d91586775e7087c5bb73。

第二阶段使用用户新提供的教材原工程，独立建立来源与适配差异，补齐GM3000实物调用。下面2.1—2.8节详细记录本次4—6学时步骤，结合第一阶段附件构成实验二完整材料；当前阶段的本地归档记录与第一阶段课程远端历史分别说明。

'''
    combined = combined.replace('## 2. 实验步骤', note+'## 2. 实验步骤', 1)
    (ROOT/(REPORT+'.md')).write_text(combined, encoding='utf-8')
    renderer = Path.home()/'.config/opencode/skills/cryptosystem-report/md2pdf.py'
    subprocess.run([sys.executable, str(renderer), str(ROOT/(PRACTICE+'.md')), str(ROOT/(PRACTICE+'.pdf'))], check=True)

    doc = Document(ROOT/'materials/报告模板.docx')
    for name in ('Heading 1', 'Heading 2', 'Heading 3'):
        if name not in doc.styles:
            doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    if 'Table Grid' not in doc.styles:
        style = doc.styles.add_style('Table Grid', WD_STYLE_TYPE.TABLE)
        borders = OxmlElement('w:tblBorders')
        for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
            element = OxmlElement('w:'+edge)
            element.set(qn('w:val'), 'single')
            element.set(qn('w:sz'), '4')
            borders.append(element)
        properties = OxmlElement('w:tblPr')
        properties.append(borders)
        style._element.append(properties)
    body = doc._element.body
    start = next(child for child in body if '实验内容、步骤与体会（附纸）' in ''.join(child.itertext()))
    deleting = False
    for child in list(body):
        if child is start:
            deleting = True
        if deleting and child.tag != qn('w:sectPr'):
            body.remove(child)
    for br in list(body.iter(qn('w:br'))):
        if br.get(qn('w:type')) == 'page':
            br.getparent().remove(br)
    for prop in list(body.iter(qn('w:pageBreakBefore'))):
        prop.getparent().remove(prop)
    for child in list(body)[::-1]:
        if child.tag == qn('w:sectPr'):
            continue
        if child.tag == qn('w:p') and not ''.join(child.itertext()).strip():
            body.remove(child)
        else:
            break
    replacements = {
        '课程：':'课程：密码系统设计                 班级：2413',
        '姓名：':'姓名：蔡贸俊                       学号：20241328',
        '成绩：':'成绩：        指导教师：赵越    实验日期：2026-10-22、10-29',
        '实验密级：':'实验密级：无     预习程度：待本人据实填写',
        '实验时间：':'实验时间：待本人核对课节  仪器组次：无  必修/选修：必修',
        '实验序号：':'实验序号：实验二',
        '实验名称：':'实验名称：密码算法实现',
        '实验目的与要求：':'实验目的与要求：在openEuler中运行检测中心与教材SM2、SM3、SM4源码，使用GmSSL和GM3000 UKey交叉验证，分析标准与函数对应关系。两阶段验证分别为36项与59项通过；本阶段Git为真实本地归档，尚未推送课程远端。',
    }
    for paragraph in doc.paragraphs:
        for prefix, value in replacements.items():
            if paragraph.text.startswith(prefix):
                paragraph.text = value
                for run in paragraph.runs:
                    font(run, size=11)
                break
    if doc.tables:
        table = doc.tables[0]
        table.cell(1, 0).text = '计算机 / GM3000 UKey'
        table.cell(1, 1).text = 'x86_64 / openEuler WSL2 / Windows SKF'
        table.cell(1, 2).text = '各1'

    def inline(paragraph, text):
        text = re.sub(r'\[([^\]]+)\]\([^)]+\)', r'\1', text)
        for part in re.split(r'(\*\*.*?\*\*|`[^`]+`)', text):
            run = paragraph.add_run(part.strip('*`') if part.startswith(('**', '`')) else part)
            font(run)
            if part.startswith('**'):
                run.bold = True
            if part.startswith('`'):
                font(run, 'Consolas', 9)

    def paragraph(text='', code=False):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.25
        if code:
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1
            font(p.add_run(text), 'Consolas', 8)
        else:
            inline(p, text)
        return p

    doc.add_page_break()
    p = doc.add_paragraph('实验内容、步骤与体会（附纸）')
    for run in p.runs:
        font(run, '黑体', 15)
    lines = combined[combined.index('## 1. 实验内容'):].splitlines()
    index, code = 0, False
    while index < len(lines):
        line = lines[index]
        if line.startswith('```'):
            code = not code
            index += 1
            continue
        if code:
            paragraph(line, True)
        elif not line.strip():
            pass
        elif line.startswith('|'):
            rows = []
            while index < len(lines) and lines[index].startswith('|'):
                cells = [c.strip() for c in lines[index].strip('|').split('|')]
                if not all(re.fullmatch('[: -]+', cell) for cell in cells):
                    rows.append(cells)
                index += 1
            table = doc.add_table(rows=0, cols=len(rows[0]))
            table.style = 'Table Grid'
            for number, row in enumerate(rows):
                cells = table.add_row().cells
                properties = table.rows[-1]._tr.get_or_add_trPr()
                properties.append(OxmlElement('w:cantSplit'))
                if number == 0:
                    properties.append(OxmlElement('w:tblHeader'))
                for cell, text in zip(cells, row):
                    p = cell.paragraphs[0]
                    p.paragraph_format.space_after = Pt(2)
                    p.paragraph_format.line_spacing = 1.05
                    inline(p, text)
                    for run in p.runs:
                        font(run, size=9)
                        run.bold = number == 0
            paragraph()
            continue
        elif line.startswith('#'):
            level = max(1, len(line)-len(line.lstrip('#'))-1)
            p = doc.add_paragraph(style='Heading '+str(min(level, 3)))
            font(p.add_run(line.lstrip('# ')), '黑体', 14 if level == 1 else 12)
            p.paragraph_format.keep_with_next = True
        elif line.startswith('!['):
            name = re.search(r'\]\((.*?)\)', line).group(1)
            p = doc.add_paragraph()
            p.add_run().add_picture(str(ROOT/name), width=Inches(6))
            paragraph('图：实际终端对本次已保存硬件输出和三方比较结果进行复核，完整证据见logs。')
        else:
            paragraph(line)
        index += 1
    for section in doc.sections:
        footer = section.footer.paragraphs[0]
        footer.alignment = 1
        field = OxmlElement('w:fldSimple')
        field.set(qn('w:instr'), 'PAGE')
        footer._p.append(field)
    docx = ROOT/(REPORT+'.docx')
    doc.save(docx)
    word = win32com.client.DispatchEx('Word.Application')
    word.Visible, word.DisplayAlerts = False, 0
    try:
        document = word.Documents.Open(str(docx))
        document.Repaginate()
        print('Template report pages:', document.ComputeStatistics(2))
        document.SaveAs2(str(ROOT/(REPORT+'.doc')), FileFormat=0)
        document.ExportAsFixedFormat(str(ROOT/(REPORT+'.pdf')), 17)
        document.Close(False)
    finally:
        word.Quit()
    review = {}
    for base, key in ((PRACTICE, 'practice'), (REPORT, 'report')):
        pdf = pymupdf.open(ROOT/(base+'.pdf'))
        texts = [p.get_text() for p in pdf]
        assert all(len(text.strip()) > 25 for text in texts), 'Possible blank PDF page'
        assert '59' in ''.join(texts) and '蔡贸俊' in ''.join(texts)
        thumbs = []
        for page in pdf:
            pix = page.get_pixmap(matrix=pymupdf.Matrix(.7, .7))
            image = Image.frombytes('RGB', [pix.width, pix.height], pix.samples)
            image.thumbnail((340, 490))
            thumbs.append(image)
        sheet = Image.new('RGB', (340*4, 520*((len(thumbs)+3)//4)), '#ddd')
        draw = ImageDraw.Draw(sheet)
        for i, image in enumerate(thumbs):
            x, y = (i%4)*340, (i//4)*520
            draw.text((x+5, y+3), 'Page '+str(i+1), fill='black')
            sheet.paste(image, (x, y+22))
        sheet.save(ROOT/f'images/{key}_pages_review.png')
        review[key] = {'pages':len(pdf), 'page_text_lengths':[len(x) for x in texts], 'no_blank_pages':True}
        (ROOT/f'logs/{key}_pdf_text.txt').write_text('\n\n'.join(texts), encoding='utf-8')
    (ROOT/'logs/pdf_review.json').write_text(json.dumps(review, indent=2), encoding='utf-8')
    print(json.dumps(review, indent=2))


if __name__ == '__main__':
    main()
